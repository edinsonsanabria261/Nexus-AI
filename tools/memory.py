"""Extracción de texto e indexación de archivos en la memoria."""

import base64
import io
import mimetypes
import os
import tempfile
from pathlib import Path

import docx2txt
import openpyxl
from openai import OpenAI
from pypdf import PdfReader

import tools.memory as memory_module


def _get_openai_client():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return None
    try:
        return OpenAI(api_key=api_key)
    except Exception as exc:
        print(f"[file_processor] No se pudo crear cliente de OpenAI: {exc}")
        return None


def _extract_image_text(file_contents: bytes, filename: str) -> list:
    client = _get_openai_client()
    if not client:
        return []

    mime_type = mimetypes.guess_type(filename)[0] or "image/png"
    encoded = base64.b64encode(file_contents).decode("utf-8")

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "Extrae todo el texto visible de esta imagen y, si no hay texto legible, "
                                "describe con detalle el contenido visual relevante."
                            ),
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:{mime_type};base64,{encoded}",
                            },
                        },
                    ],
                }
            ],
        )
        text = (response.choices[0].message.content or "").strip()
        return [(text, 1)] if text else []
    except Exception as exc:
        print(f"[file_processor] Error leyendo imagen {filename}: {exc}")
        return []


def _extract_audio_text(file_contents: bytes, filename: str) -> list:
    client = _get_openai_client()
    if not client:
        return []

    suffix = Path(filename).suffix or ".wav"
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(file_contents)
            tmp_path = tmp.name

        try:
            with open(tmp_path, "rb") as audio_file:
                response = client.audio.transcriptions.create(
                    model="whisper-1",
                    file=(filename, audio_file),
                )
            text = (getattr(response, "text", "") or "").strip()
            return [(text, 1)] if text else []
        finally:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass
    except Exception as exc:
        print(f"[file_processor] Error transcribiendo audio {filename}: {exc}")
        return []


def _extract(file_contents: bytes, filename: str) -> list:
    stream = io.BytesIO(file_contents)
    name = filename.lower()
    pages = []

    if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")):
        return _extract_image_text(file_contents, filename)

    if name.endswith((".mp3", ".wav", ".m4a", ".ogg", ".aac", ".flac")):
        return _extract_audio_text(file_contents, filename)

    if name.endswith(".pdf"):
        reader = PdfReader(stream)
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append((text, number))
        return pages

    if name.endswith(".docx") or name.endswith(".doc"):
        text = docx2txt.process(stream) or ""
        return [(text, 1)] if text.strip() else []

    if name.endswith(".xlsx"):
        workbook = openpyxl.load_workbook(stream, data_only=True)
        parts = []
        for sheet_name in workbook.sheetnames:
            parts.append(f"--- Hoja: {sheet_name} ---")
            for row in workbook[sheet_name].iter_rows(values_only=True):
                line = "\t".join(str(cell) for cell in row if cell is not None)
                if line.strip():
                    parts.append(line)
        text = "\n".join(parts)
        return [(text, 1)] if text.strip() else []

    text = ""
    for encoding in ("utf-8", "latin-1", "iso-8859-15", "cp1252"):
        try:
            text = file_contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = file_contents.decode("utf-8", errors="ignore")
    return [(text, 1)] if text.strip() else []


def _insert_chunk(client, content: str, filename: str, page_num: int) -> bool:
    payload = {
        "content": content,
        "source": "file_upload",
        "tags": f"file_upload:{filename}",
        "filename": filename,
        "page_number": int(page_num),
    }
    try:
        client.table("knowledge").insert(payload).execute()
        return True
    except Exception as exc:
        print(f"[file_processor] No se pudo guardar un bloque: {exc}")
        return False


def process_and_index_file(file_contents: bytes, filename: str) -> int:
    try:
        pages = _extract(file_contents, filename)
    except Exception as exc:
        print(f"[file_processor] Error leyendo {filename}: {exc}")
        return 0
    if not pages:
        return 0

    client = memory_module.get_supabase()
    if not client:
        print("[file_processor] Sin Supabase: el texto se leyó, pero no hay dónde guardarlo.")
        return 0

    saved = 0
    for content, page_num in pages:
        if not content or not str(content).strip():
            continue
        words = str(content).split()
        step = 450
        size = 600
        for start in range(0, len(words), step):
            chunk = " ".join(words[start : start + size]).strip()
            if len(chunk) < 20:
                continue
            formatted = f"[Archivo: {filename} | Pág: {page_num}]\n{chunk}"
            if _insert_chunk(client, formatted, filename, page_num):
                saved += 1
    return saved
