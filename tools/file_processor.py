"""Extracción de texto e indexación de archivos en la memoria."""

import io

import docx2txt
import openpyxl
from pypdf import PdfReader

import tools.memory as memory_module


def _extract(file_contents: bytes, filename: str) -> list:
    stream = io.BytesIO(file_contents)
    name = filename.lower()
    pages = []

    if name.endswith(".pdf"):
        reader = PdfReader(stream)
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append((text, number))
        return pages

    if name.endswith(".docx") or name.endswith(".doc"):
        text = docx2txt.process(stream) or ""
        return [(text, 1)] if text.strip() else []

    if name.endswith(".xlsx"):
        workbook = openpyxl.load_workbook(stream, data_only=True)
        parts = []
        for sheet_name in workbook.sheetnames:
            parts.append(f"--- Hoja: {sheet_name} ---")
            for row in workbook[sheet_name].iter_rows(values_only=True):
                line = "\t".join(str(cell) for cell in row if cell is not None)
                if line.strip():
                    parts.append(line)
        text = "\n".join(parts)
        return [(text, 1)] if text.strip() else []

    text = ""
    for encoding in ("utf-8", "latin-1", "iso-8859-15", "cp1252"):
        try:
            text = file_contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = file_contents.decode("utf-8", errors="ignore")
    return [(text, 1)] if text.strip() else []


def _insert_chunk(client, content: str, filename: str, page_num: int) -> bool:
    embedding = memory_module.obtener_embedding_externo(content)
    full = {
        "content": content,
        "source": "file_upload",
        "tags": f"file_upload:{filename}",
        "filename": filename,
        "page_number": int(page_num),
    }
    if embedding:
        full["embedding"] = embedding
    minimal = {
        "content": content,
        "source": "file_upload",
        "tags": f"file_upload:{filename}",
    }
    last_error = None
    for payload in (full, minimal):
        try:
            client.table("knowledge").insert(payload).execute()
            return True
        except Exception as exc:
            last_error = exc
    print(f"[file_processor] No se pudo guardar un bloque: {last_error}")
    return False


def process_and_index_file(file_contents: bytes, filename: str) -> int:
    try:
        pages = _extract(file_contents, filename)
    except Exception as exc:
        print(f"[file_processor] Error leyendo {filename}: {exc}")
        return 0
    if not pages:
        return 0

    client = memory_module.get_supabase()
    if not client:
        print("[file_processor] Sin Supabase: el texto se leyó, pero no hay dónde guardarlo.")
        return 0

    saved = 0
    for content, page_num in pages:
        words = content.split()
        step = 450
        size = 600
        for start in range(0, len(words), step):
            chunk = " ".join(words[start : start + size]).strip()
            if len(chunk) < 20:
                continue
            formatted = f"[Archivo: {filename} | Pág: {page_num}]\n{chunk}"
            if _insert_chunk(client, formatted, filename, page_num):
                saved += 1
    return saved
